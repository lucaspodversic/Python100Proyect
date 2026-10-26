"""Exportación de resultados a CSV y Excel (spreadsheet).

Opcionalmente compara contra nuestros precios (``propios``: sku -> precio)
agregando la diferencia porcentual, que es lo que el equipo comercial usa.
"""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .modelos import Producto


def _filas(productos: list[Producto], propios: dict[str, Decimal] | None) -> tuple[list[str], list[list[object]]]:
    """Construye encabezados y filas comunes a ambos formatos."""
    # Unión de todas las columnas extra de todos los competidores, ordenada para un encabezado
    # estable; si un producto no tiene una de esas columnas, la celda queda vacía.
    extras = sorted({k for p in productos for k in p.extras})
    cab = ["competidor", "sku", "titulo", "precio", "url", *extras]
    if propios is not None:
        cab += ["precio_propio", "diferencia_%"]
    filas: list[list[object]] = []
    for p in productos:
        fila: list[object] = [p.competidor, p.sku or "", p.titulo, p.precio, p.url or ""]
        fila += [p.extras.get(k, "") for k in extras]
        if propios is not None:
            # Se compara por SKU; sin SKU (o SKU no presente en ``propios``) la diferencia queda vacía.
            nuestro = propios.get(p.sku or "")
            dif = None
            # Caso borde: precio "Consultar" (None) o precio propio 0 → no se calcula (evita división por cero).
            if nuestro and p.precio is not None:
                # Positivo => la competencia es más cara que nosotros.
                dif = round(float((p.precio - nuestro) / nuestro * 100), 2)
            fila += [nuestro, dif]
        filas.append(fila)
    return cab, filas


def exportar_csv(productos: list[Producto], destino: Path, propios: dict[str, Decimal] | None = None) -> Path:
    """Escribe un CSV UTF-8 con BOM (Excel lo abre con acentos correctos)."""
    cab, filas = _filas(productos, propios)
    # newline="" es lo que pide el módulo csv para no duplicar saltos de línea en Windows.
    with destino.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cab)
        # None se escribe como celda vacía (csv escribiría la cadena vacía igual, pero lo hacemos explícito).
        w.writerows(["" if v is None else v for v in fila] for fila in filas)
    return destino


def exportar_xlsx(productos: list[Producto], destino: Path, propios: dict[str, Decimal] | None = None) -> Path:
    """Escribe un .xlsx con encabezado en negrita, filtros y formato moneda."""
    cab, filas = _filas(productos, propios)
    wb = Workbook()
    ws = wb.active
    ws.title = "Competencia"
    ws.append(cab)
    # Fila 1 = encabezado: texto blanco en negrita sobre azul, para distinguirla de los datos.
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="305496")
    for fila in filas:
        # openpyxl no serializa Decimal de forma nativa en todas las versiones: se convierte a float
        # sólo al escribir el Excel (los cálculos ya se hicieron con Decimal).
        ws.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    # openpyxl indexa columnas desde 1; formato de miles con 2 decimales en la columna de precio.
    col_precio = cab.index("precio") + 1
    for (celda,) in ws.iter_rows(min_row=2, min_col=col_precio, max_col=col_precio):
        celda.number_format = "#,##0.00"
    # Autofiltro sobre todo el rango usado y encabezado fijo al hacer scroll (freeze_panes "A2").
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(destino)
    return destino
