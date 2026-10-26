"""Exportación a .xlsx: una hoja por tabla, encabezado en negrita, panel congelado y anchos ajustados."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extraer import Tabla


def exportar_excel(tablas: list[Tabla], ruta: Path, formato_numero: str = "#,##0.00") -> Path:
    """Escribe un libro con una hoja por tabla; ValueError si la lista está vacía."""
    if not tablas:
        # Un .xlsx sin hojas es inválido para Excel: mejor fallar con un mensaje claro.
        raise ValueError("No hay tablas para exportar")
    wb = Workbook()
    # El libro nuevo trae una hoja "Sheet" vacía: se quita para que sólo queden las tablas.
    wb.remove(wb.active)
    usados: set[str] = set()
    for t in tablas:
        # Nombre que permite ubicar la tabla en el PDF: p<página>_t<n.º>.
        nombre = f"p{t.pagina}_t{t.indice + 1}"
        while nombre in usados:      # nombres de hoja únicos (máx. 31 caracteres en Excel)
            nombre += "_"
        usados.add(nombre)
        # Excel limita los nombres de hoja a 31 caracteres (sólo se alcanza con muchísimos duplicados).
        ws = wb.create_sheet(nombre[:31])
        for fila in t.filas:
            ws.append(fila)
        # Encabezado destacado: negrita + fondo celeste claro.
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        # Congela la fila 1: el encabezado queda visible al desplazarse.
        ws.freeze_panes = "A2"
        # Formato numérico sólo en celdas float (las de texto no se tocan).
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                if isinstance(c.value, float):
                    c.number_format = formato_numero
        # Ancho de columna ≈ contenido más largo + margen, con tope de 60 para textos largos.
        for j, col in enumerate(ws.columns, 1):
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(j)].width = min(60, ancho + 2)
    # Crea la carpeta destino si no existe (ej. "salidas/2026/informe.xlsx").
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return Path(ruta)
