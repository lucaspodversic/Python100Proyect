"""Tests con PDFs generados al vuelo (tests/pdf_fixture.py), sin archivos binarios versionados."""
import pytest
from openpyxl import load_workbook

from pdftablas import Tabla, convertir_numero, exportar_excel, extraer_tablas, limpiar, unir_continuaciones
from pdftablas.cli import ejecutar
from pdftablas.extraer import parsear_paginas
from pdf_fixture import crear_pdf

# Encabezado + 3 filas, con números en formato latino.
FILAS = [("Producto", "Cantidad", "Precio"), ("Manzana", "10", "1.250,50"), ("Pera", "3", "980,00"), ("Uva roja", "7", "2.100,75")]


def pagina(filas, titulo=None):
    """Ubica las celdas en una grilla sin bordes: columnas cada 150 pt, filas cada 20 pt."""
    # El título opcional va arriba (y=750), separado de la tabla: limpiar() debe descartarlo.
    celdas = [(72, 750, titulo)] if titulo else []
    for i, f in enumerate(filas):
        celdas += [(72 + 150 * j, 700 - 20 * i, c) for j, c in enumerate(f)]
    return celdas


@pytest.fixture
def pdf(tmp_path):
    """PDF de 3 páginas: tabla con título, continuación con encabezado repetido, y página sin tabla."""
    ruta = tmp_path / "informe.pdf"
    ruta.write_bytes(crear_pdf([pagina(FILAS, "Informe de ventas"),
                                # pdfplumber (estrategia texto) pide >= 3 palabras alineadas por columna
                                pagina([FILAS[0], ("Lima", "1", "9,00"), ("Kiwi", "2", "50,25")]),
                                [(72, 700, "Sin tablas aqui")]]))
    return ruta


def test_paginas():
    # Rango + página suelta, ordenadas y sin duplicados.
    assert parsear_paginas("1,3-4", 5) == [1, 3, 4]
    assert parsear_paginas(None, 2) == [1, 2]
    # Caso borde: rango que excede el PDF.
    with pytest.raises(ValueError):
        parsear_paginas("2-9", 3)


# Ambos formatos decimales, porcentaje, negativo, código alfanumérico (no se convierte) y vacío.
@pytest.mark.parametrize("texto,dec,esperado", [("1.250,50", ",", 1250.5), ("1,250.50", ".", 1250.5),
                                                ("15%", ",", 0.15), ("-3", ",", -3.0), ("A-12", ",", "A-12"), ("", ",", "")])
def test_convertir_numero(texto, dec, esperado):
    assert convertir_numero(texto, dec) == esperado


def test_extraer_y_limpiar(pdf):
    crudas = extraer_tablas(pdf, "1", "texto")
    assert len(crudas) == 1 and crudas[0].pagina == 1
    t = limpiar(crudas[0])
    # Sin el título, encabezado como texto, números convertidos y "Uva roja" en una celda.
    assert t.filas == [["Producto", "Cantidad", "Precio"], ["Manzana", 10.0, 1250.5], ["Pera", 3.0, 980.0],
                       ["Uva roja", 7.0, 2100.75]]
    with pytest.raises(ValueError):
        extraer_tablas(pdf, estrategia="magia")


def test_limpiar_columnas_vacias():
    # Columna central y fila central vacías se eliminan; "x\n y" y " 2 " se normalizan.
    # (`None or ""` simplemente vale "": representa la celda que pdfplumber devolvería como None.)
    t = limpiar(Tabla(1, 0, [["A", None or "", "B"], ["", "", ""], ["x\n y", "", " 2 "]]), numeros=False)
    assert t.filas == [["A", "B"], ["x y", "2"]]


def test_unir_continuaciones(pdf):
    tablas = [limpiar(t) for t in extraer_tablas(pdf, "1-2", "texto")]
    unidas = unir_continuaciones(tablas)
    # 1 encabezado + 3 filas (pág. 1) + 2 filas (pág. 2, sin repetir el encabezado) = 6.
    assert len(unidas) == 1 and unidas[0].filas[-1] == ["Kiwi", 2.0, 50.25] and len(unidas[0].filas) == 6
    # Continuación sin encabezado (primera fila numérica): se agrega completa.
    sin_enc = unir_continuaciones([Tabla(1, 0, [["A", "B"], ["x", 1.0]]), Tabla(2, 0, [[2.0, 3.0]])])
    assert sin_enc[0].filas[-1] == [2.0, 3.0]


def test_excel(tmp_path):
    # Dos tablas con la misma página e índice → la segunda hoja recibe "_" para ser única.
    ruta = exportar_excel([Tabla(1, 0, [["N", "V"], ["a", 1234.5]]), Tabla(1, 0, [["X"], ["y"]])], tmp_path / "o.xlsx")
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["p1_t1", "p1_t1_"]
    ws = wb["p1_t1"]
    assert ws["B2"].value == 1234.5 and ws["B2"].number_format == "#,##0.00" and ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    with pytest.raises(ValueError):
        exportar_excel([], tmp_path / "x.xlsx")


def test_cli(pdf, tmp_path, capsys):
    # La carpeta "out" no existe: exportar_excel debe crearla.
    salida = tmp_path / "out" / "t.xlsx"
    assert ejecutar([str(pdf), str(salida), "--estrategia", "texto", "--unir"]) == 0
    # 3 + 2 filas de datos unidas en una sola hoja.
    assert "5 fila(s) × 3 col." in capsys.readouterr().out
    assert load_workbook(salida).sheetnames == ["p1_t1"]
    # Página 3 sin tabla → exit 2; PDF inexistente o corrupto → exit 1.
    assert ejecutar([str(pdf), str(salida), "--paginas", "3", "--estrategia", "texto"]) == 2
    assert ejecutar([str(tmp_path / "no.pdf"), str(salida)]) == 1
    roto = tmp_path / "roto.pdf"; roto.write_bytes(b"no soy un pdf")
    assert ejecutar([str(roto), str(salida)]) == 1
